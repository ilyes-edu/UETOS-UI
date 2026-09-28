"""Export a timetable version to Excel (.xlsx) and PDF – coloured by subject, Arabic / English.

Sections: full view of all classes, full view of all teachers, one timetable per class, one per teacher."""
import io
import os
import pandas as pd

import colors
import i18n
from i18n import t

HERE = os.path.dirname(os.path.abspath(__file__))
DAYS, SLOTS = 5, 7
OFF = {(2, s) for s in range(4, 7)}          # Tuesday afternoon


# ------------------------------------------------------------------ data
def _class_key(c):
    import re
    return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", str(c))]


def views(df):
    """-> (class_cells, teacher_cells): {row: {(d, s): (title, sub, subject)}}  (remediation shown in every class)."""
    df = df.copy()
    if "Classes" not in df.columns:
        df["Classes"] = ""
    df["Classes"] = df["Classes"].fillna("").astype(str)
    cc, tc = {}, {}
    for r in df.itertuples():
        rem = str(r.Class).startswith("REM:") or r.Subject == "REMEDIAL"
        tchs = [x for x in str(r.Teachers).split(", ") if x]
        rows = [z.strip() for z in r.Classes.split(",") if z.strip()] if rem else [r.Class]
        for c in rows:
            cell = cc.setdefault(c, {})
            title, sub = i18n.subj_short("REMEDIAL" if rem else r.Subject), i18n.teachers(tchs)
            if (r.Day, r.Slot) in cell:                       # grouped remediation etc.
                o = cell[(r.Day, r.Slot)]
                title, sub = o[0], o[1] + " / " + sub
            cell[(r.Day, r.Slot)] = (title, sub, r.Subject)
        for tch in tchs:
            n = len(rows)
            tc.setdefault(tch, {})[(r.Day, r.Slot)] = (
                i18n.subj_short("REMEDIAL") if rem else i18n.cls(r.Class),
                t("n_classes", n=n) if rem else i18n.subj_short(r.Subject), r.Subject)
    classes = sorted(cc, key=_class_key)
    teachers = sorted(tc, key=_class_key)
    return {c: cc[c] for c in classes}, {x: tc[x] for x in teachers}


def _rowlab(is_class):
    return t("class") if is_class is True else (t("rm_room") if is_class == "room" else t("teacher"))


def room_view(occ, caps):
    """Occupancy -> {room: {(d, s): (classes, short subject, subject)}} in room order (for the full-grid layout)."""
    import rooms as rms
    names = rms.room_names(caps)
    out = {r: {} for rt in names for r in names[rt]}
    for r in occ.itertuples():
        if r.Room:
            cls_ = ", ".join(i18n.cls(c.strip()) for c in str(r.Class).split(",")[:3])
            out.setdefault(r.Room, {})[(int(r.Day), int(r.Slot))] = (cls_, i18n.subj_short(r.Subject), r.Subject)
    return out


def _hours(cells):
    return sum(1 for (d, s), v in cells.items() if v[2] not in ("REMEDIAL", "RECEPTION"))


def _add_reception(tc, reception):
    """Reception hours as a teacher-only cell (not counted in the hours)."""
    lab = t("rec_title").replace("🤝", "").strip()
    for tch, (d, s) in (reception or {}).items():
        if tch in tc and (int(d), int(s)) not in tc[tch]:
            tc[tch][(int(d), int(s))] = (lab, lab, "RECEPTION")


# ------------------------------------------------------------------ XLSX
def to_xlsx(df, title="", sections=("full_class", "full_teacher", "class", "teacher"), rooms=None, reception=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cc, tc = views(df)
    _add_reception(tc, reception)
    wb = Workbook(); wb.remove(wb.active)
    rtl = i18n.is_ar()
    thin = Side(style="thin", color="BBBBBB"); thick = Side(style="medium", color="777777")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hdr_fill = PatternFill("solid", fgColor="E8EAF0"); off_fill = PatternFill("solid", fgColor="EDEDED")
    fill = lambda subj: PatternFill("solid", fgColor=colors.light(subj).lstrip("#"))

    def sheet_name(x, used):
        base = "".join(ch for ch in str(x) if ch not in '[]:*?/\\')[:28] or "Sheet"
        n, k = base, 2
        while n in used:
            n, k = f"{base[:25]}_{k}", k + 1
        used.add(n); return n
    used = set()

    def full_sheet(name, data, is_class):
        ws = wb.create_sheet(sheet_name(name, used)); ws.sheet_view.rightToLeft = rtl
        ws.cell(1, 1, title).font = Font(bold=True, size=12)
        ws.cell(2, 1, _rowlab(is_class)).font = Font(bold=True)
        ws.cell(2, 1).fill = hdr_fill; ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
        ws.cell(2, 2, t("tot_hours")).font = Font(bold=True); ws.cell(2, 2).fill = hdr_fill
        ws.merge_cells(start_row=2, start_column=2, end_row=3, end_column=2)
        for d in range(DAYS):
            c0 = 3 + d * SLOTS
            ws.cell(2, c0, i18n.day(d)).font = Font(bold=True)
            ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + SLOTS - 1)
            for s in range(SLOTS):
                cl = ws.cell(3, c0 + s, s + 1); cl.alignment = center; cl.fill = hdr_fill; cl.font = Font(size=8, bold=True)
            for s in range(SLOTS):
                ws.cell(2, c0 + s).fill = hdr_fill; ws.cell(2, c0 + s).alignment = center
        for i, (key, cells) in enumerate(data.items()):
            r = 4 + i
            ws.cell(r, 1, i18n.cls(key) if is_class is True else key).font = Font(bold=True)
            ws.cell(r, 2, _hours(cells)).alignment = center
            for d in range(DAYS):
                for s in range(SLOTS):
                    cl = ws.cell(r, 3 + d * SLOTS + s)
                    cl.alignment = center; cl.font = Font(size=7)
                    cl.border = Border(left=thin, right=thick if s == SLOTS - 1 else thin, top=thin, bottom=thin)
                    if (d, s) in cells:
                        ti, sub, subj = cells[(d, s)]
                        cl.value = f"{ti}\n{sub}"; cl.fill = fill(subj)
                    elif (d, s) in OFF:
                        cl.fill = off_fill
            ws.row_dimensions[r].height = 30
        ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 6
        for col in range(3, 3 + DAYS * SLOTS):
            ws.column_dimensions[get_column_letter(col)].width = 8.5
        ws.freeze_panes = "C4"
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True

    def one_sheet(key, cells, is_class):
        ws = wb.create_sheet(sheet_name(i18n.cls(key) if is_class else key, used)); ws.sheet_view.rightToLeft = rtl
        ws.cell(1, 1, (f"{title} — " if title else "") + (i18n.cls(key) if is_class else key)
                + f" — {t('tot_h', h=_hours(cells))}").font = Font(bold=True, size=13)
        ws.cell(2, 1, "").fill = hdr_fill
        for d in range(DAYS):
            cl = ws.cell(2, 2 + d, i18n.day(d)); cl.font = Font(bold=True); cl.fill = hdr_fill; cl.alignment = center
        for s in range(SLOTS):
            cl = ws.cell(3 + s, 1, i18n.slot_label(s)); cl.font = Font(bold=True); cl.fill = hdr_fill; cl.alignment = center
            ws.row_dimensions[3 + s].height = 38
            for d in range(DAYS):
                cl = ws.cell(3 + s, 2 + d); cl.alignment = center; cl.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if (d, s) in cells:
                    ti, sub, subj = cells[(d, s)]
                    cl.value = f"{ti}\n{sub}"; cl.fill = fill(subj); cl.font = Font(size=10)
                elif (d, s) in OFF:
                    cl.fill = off_fill
        ws.column_dimensions["A"].width = 18
        for d in range(DAYS):
            ws.column_dimensions[get_column_letter(2 + d)].width = 20
        ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True

    if "full_class" in sections: full_sheet(t("lay_full_class").replace("🗂", "").strip(), cc, True)
    if "full_teacher" in sections: full_sheet(t("lay_full_teacher").replace("🗂", "").strip(), tc, False)
    if "rooms" in sections and rooms: full_sheet(t("rm_view").replace("🏫", "").strip(), rooms, "room")
    if "class" in sections:
        for k, v in cc.items(): one_sheet(k, v, True)
    if "teacher" in sections:
        for k, v in tc.items(): one_sheet(k, v, False)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ------------------------------------------------------------------ PDF
_FONT_OK = False


def _fonts():
    global _FONT_OK
    if _FONT_OK:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("DV", os.path.join(HERE, "fonts", "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DVB", os.path.join(HERE, "fonts", "DejaVuSans-Bold.ttf")))
    _FONT_OK = True


def _ar(txt):
    """Shape + reorder Arabic so reportlab draws it correctly."""
    txt = str(txt)
    if not any("\u0600" <= ch <= "\u06ff" for ch in txt):
        return txt
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(txt))
    except Exception:
        return txt


def to_pdf(df, title="", sections=("full_class", "full_teacher", "class", "teacher"), rooms=None, reception=None):
    from reportlab.lib import colors as rc
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _fonts()
    cc, tc = views(df)
    _add_reception(tc, reception)
    rtl = i18n.is_ar()
    buf = io.BytesIO()
    W, Hh = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=30, bottomMargin=18,
                            title=title)

    def _header(canvas, _doc):          # school name — version on every page
        if not title:
            return
        canvas.saveState(); canvas.setFont("DV", 8); canvas.setFillGray(0.35)
        (canvas.drawRightString(W - 18, Hh - 16, _ar(title)) if rtl else canvas.drawString(18, Hh - 16, _ar(title)))
        canvas.restoreState()
    h1 = ParagraphStyle("h1", fontName="DVB", fontSize=13, alignment=2 if rtl else 0, spaceAfter=6)
    cellst = lambda size, bold=False: ParagraphStyle("c", fontName="DVB" if bold else "DV", fontSize=size,
                                                     leading=size + 1.2, alignment=1)
    P = lambda txt, size=6, bold=False: Paragraph(_ar(txt).replace("\n", "<br/>"), cellst(size, bold))
    story = []

    def full_table(name, data, is_class):
        days = list(range(DAYS))[::-1] if rtl else list(range(DAYS))
        slots_order = lambda: list(range(SLOTS))[::-1] if rtl else list(range(SLOTS))
        head1, head2 = [], []
        for d in days:
            head1 += [P(i18n.day(d), 7, True)] + [""] * (SLOTS - 1)
            head2 += [P(str(s + 1), 6, True) for s in slots_order()]
        lab = [P(_rowlab(is_class), 7, True), P(t("tot_hours"), 5.5, True)]
        head1 = (head1 + lab[::-1]) if rtl else (lab + head1)
        head2 = (head2 + ["", ""]) if rtl else (["", ""] + head2)
        rows = [head1, head2]
        styles = []
        per_page = 22
        items = list(data.items())
        chunks = [items[i:i + per_page] for i in range(0, len(items), per_page)] or [[]]
        for ci, chunk in enumerate(chunks):
            rows = [head1, head2]; styles = []
            for ri, (key, cells) in enumerate(chunk):
                line = []
                for d in days:
                    for s in slots_order():
                        if (d, s) in cells:
                            ti, sub, subj = cells[(d, s)]
                            line.append(P(f"{ti}\n{sub}", 4.6))
                            col = len(line) - 1 + (0 if rtl else 2)
                            styles.append(("BACKGROUND", (col, ri + 2), (col, ri + 2), rc.HexColor(colors.light(subj))))
                        else:
                            line.append("")
                            if (d, s) in OFF:
                                col = len(line) - 1 + (0 if rtl else 2)
                                styles.append(("BACKGROUND", (col, ri + 2), (col, ri + 2), rc.HexColor("#ececec")))
                lbl = [P(i18n.cls(key) if is_class is True else key, 6.5, True), P(str(_hours(cells)), 6)]
                rows.append((line + lbl[::-1]) if rtl else (lbl + line))
            cw = (W - 36 - 70) / (DAYS * SLOTS)
            widths = ([cw] * DAYS * SLOTS + [22, 48]) if rtl else ([48, 22] + [cw] * DAYS * SLOTS)
            tbl = Table(rows, colWidths=widths, rowHeights=[13, 10] + [min(22, (Hh - 90) / max(1, len(chunk)))] * len(chunk),
                        repeatRows=2)
            base = [("GRID", (0, 0), (-1, -1), 0.25, rc.HexColor("#bbbbbb")),
                    ("BACKGROUND", (0, 0), (-1, 1), rc.HexColor("#e8eaf0")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0.5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0.5), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]
            off = 0 if rtl else 2
            for k in range(DAYS):
                c0 = off + k * SLOTS
                base.append(("SPAN", (c0, 0), (c0 + SLOTS - 1, 0)))
                base.append(("LINEAFTER" if not rtl else "LINEBEFORE", (c0 + (SLOTS - 1 if not rtl else 0), 0),
                             (c0 + (SLOTS - 1 if not rtl else 0), -1), 1.2, rc.HexColor("#666666")))
            lab0 = DAYS * SLOTS if rtl else 0
            base += [("SPAN", (lab0, 0), (lab0, 1)), ("SPAN", (lab0 + 1, 0), (lab0 + 1, 1))]
            tbl.setStyle(TableStyle(base + styles))
            story.append(Paragraph(_ar(f"{title} — {name}" + (f" ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else "")), h1))
            story.append(tbl); story.append(PageBreak())

    def one_table(key, cells, is_class):
        days = list(range(DAYS))[::-1] if rtl else list(range(DAYS))
        head = [P(i18n.day(d), 10, True) for d in days]
        head = (head + [""]) if rtl else ([""] + head)
        rows, styles = [head], []
        for s in range(SLOTS):
            line = []
            for d in days:
                if (d, s) in cells:
                    ti, sub, subj = cells[(d, s)]
                    line.append(P(f"{ti}\n{sub}", 9))
                    col = len(line) - 1 + (0 if rtl else 1)
                    styles.append(("BACKGROUND", (col, s + 1), (col, s + 1), rc.HexColor(colors.light(subj))))
                else:
                    line.append("")
                    if (d, s) in OFF:
                        col = len(line) - 1 + (0 if rtl else 1)
                        styles.append(("BACKGROUND", (col, s + 1), (col, s + 1), rc.HexColor("#ececec")))
            lbl = P(i18n.slot_label(s), 8, True)
            rows.append((line + [lbl]) if rtl else ([lbl] + line))
        cw = (W - 36 - 90) / DAYS
        widths = ([cw] * DAYS + [90]) if rtl else ([90] + [cw] * DAYS)
        tbl = Table(rows, colWidths=widths, rowHeights=[20] + [58] * SLOTS)
        tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, rc.HexColor("#999999")),
                                 ("BACKGROUND", (0, 0), (-1, 0), rc.HexColor("#e8eaf0")),
                                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE")] + styles))
        name = i18n.cls(key) if is_class else key
        story.append(Paragraph(_ar(f"{'🎒 ' if False else ''}{name} — {t('tot_h', h=_hours(cells))}"), h1))
        story.append(tbl); story.append(PageBreak())

    if "full_class" in sections: full_table(t("lay_full_class").replace("🗂", "").strip(), cc, True)
    if "full_teacher" in sections: full_table(t("lay_full_teacher").replace("🗂", "").strip(), tc, False)
    if "rooms" in sections and rooms: full_table(t("rm_view").replace("🏫", "").strip(), rooms, "room")
    if "class" in sections:
        for k, v in cc.items(): one_table(k, v, True)
    if "teacher" in sections:
        for k, v in tc.items(): one_table(k, v, False)
    if story and isinstance(story[-1], PageBreak):
        story.pop()
    if not story:
        story = [Spacer(1, 10)]
    doc.build(story, onFirstPage=_header, onLaterPages=_header)
    return buf.getvalue()
