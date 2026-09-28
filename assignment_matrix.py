"""Import / export of the school 'teacher assignment matrix' (جدول توزيع الحصص).

Layout (as used in Algerian middle schools, sheet displayed right-to-left):
    row  "الحجم الساعي" : weekly total per post
    row  "المناصب"      : post names  (عربية1, رياضيات2, فيزياء3, اعلام آلي, رياضة 1 …)
    row  "اللقب و الاسم" : teacher names (optional)
    rows  رقم | الافواج (class e.g. 4م1) | حجم ساعي | hours per post …
    row  "الاستدراك"     : remedial hours per post (optional)
    row  "المجموع"       : totals
The parser is tolerant: it locates the post row and the class rows by content, not by position.
"""
import io
import re
import pandas as pd
from i18n import t

# ---------------------------------------------------------------- subject names
def normalize(s):
    s = str(s)
    s = re.sub(r"[\d\s\u0640\-_.()]", "", s)          # digits, spaces, tatweel, punctuation
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ى", "ي")
    return s.lower()

# normalized Arabic base name -> candidate curriculum codes (first one present in the curriculum wins)
SUBJECT_ALIASES = {
    "عربيه": ["ARABIC"], "لغهعربيه": ["ARABIC"],
    "اسلاميه": ["ISLAMIC"], "تربيهاسلاميه": ["ISLAMIC"],
    "رياضيات": ["MATH", "MATHS"],
    "فرنسيه": ["FRENCH"], "لغهفرنسيه": ["FRENCH"],
    "انجليزيه": ["ENGLISH"], "انكليزيه": ["ENGLISH"], "لغهانجليزيه": ["ENGLISH"],
    "فيزياء": ["PHYSICS", "PHYS"], "علومفيزيائيه": ["PHYSICS", "PHYS"],
    "علوم": ["SCIENCE", "SVT"], "علومطبيعيه": ["SCIENCE", "SVT"],
    "اجتماعيات": ["HISTGEO", "HG"], "تاريخوجغرافيا": ["HISTGEO", "HG"],
    "اعلامالي": ["INFO", "CS"], "اعلاميالي": ["INFO", "CS"], "معلوماتيه": ["INFO", "CS"],
    "موسيقي": ["MUSIC"], "تربيهموسيقيه": ["MUSIC"],
    "رياضه": ["SPORT", "PE", "EPS"], "تربيهبدنيه": ["SPORT", "PE", "EPS"],
    "تشكيليه": ["ART"], "تربيهتشكيليه": ["ART"], "رسم": ["ART"],
    "مدنيه": ["CIVIC", "CIVICS"], "تربيهمدنيه": ["CIVIC", "CIVICS"],
    "تاريخ": ["HIST"], "جغرافيا": ["GEO"], "فنيه": ["ART_MUSIC"], "تربيهفنيه": ["ART_MUSIC"],
    "امازيغيه": ["AMAZIGH", "TAMAZIGHT"], "تكنولوجيا": ["TECH"],
}
CODE_TO_ARABIC = {"ARABIC": "عربية", "ISLAMIC": "اسلامية", "MATH": "رياضيات", "FRENCH": "فرنسية",
                  "ENGLISH": "انجليزية", "PHYSICS": "فيزياء", "SCIENCE": "علوم", "HISTGEO": "اجتماعيات",
                  "INFO": "اعلام آلي", "MUSIC": "موسيقى", "SPORT": "رياضة", "PE": "رياضة", "ART": "تربية تشكيلية",
                  "CIVIC": "مدنية", "AMAZIGH": "امازيغية", "TECH": "تكنولوجيا",
                  "PHYS": "فيزياء", "SVT": "علوم", "HG": "اجتماعيات", "HIST": "تاريخ", "GEO": "جغرافيا",
                  "CIVICS": "مدنية", "ART_MUSIC": "فنية", "CS": "اعلام آلي", "EPS": "رياضة", "MATHS": "رياضيات"}

CLASS_RE = re.compile(r"^\s*(\d)\s*(?:م|AM|am)\s*(\d+)\s*$")


def resolve_subject(header, curriculum_codes):
    """'فيزياء2' -> 'PHYSICS' (or None if the header is not a subject)."""
    raw = str(header).strip()
    if not raw or raw.lower() == "nan":
        return None
    latin = re.sub(r"[\d\s_]+$", "", raw).upper()          # post names may already be codes (MATH1)
    if latin in curriculum_codes:
        return latin
    cands = SUBJECT_ALIASES.get(normalize(raw))
    if not cands:
        return None
    for c in cands:
        if c in curriculum_codes:
            return c
    return cands[0]


def class_to_id(label):
    """'4م1' / '4AM1' -> '4AM1'."""
    m = CLASS_RE.match(str(label))
    return f"{m.group(1)}AM{m.group(2)}" if m else None


def _num(v):
    try:
        f = float(str(v).strip())
        return f if f == f else 0.0
    except ValueError:
        return 0.0


# ---------------------------------------------------------------- parsing
def read_any(file):
    name = getattr(file, "name", str(file)).lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(file, header=None, dtype=object)
    return pd.read_csv(file, header=None, dtype=object)


def parse_assignment_matrix(raw, curriculum_codes, known_classes=None):
    """raw: DataFrame read with header=None.  Returns dict with
       assignment (Teacher, Class, Subject, Hours, Post), teachers, classes, remedial, warnings."""
    warnings = []
    curriculum_codes = set(curriculum_codes)
    grid = raw.fillna("").astype(str).map(lambda x: x.strip())

    # 1. locate the posts row: the row with most cells that resolve to a subject
    best, post_row = 0, None
    for i in range(min(len(grid), 30)):
        n = sum(resolve_subject(v, curriculum_codes | {"__"}) is not None for v in grid.iloc[i])
        if n > best:
            best, post_row = n, i
    if post_row is None or best < 2:
        raise ValueError(t("w_no_posts"))
    posts = {}                                           # col -> (post name, subject code)
    for j, v in enumerate(grid.iloc[post_row]):
        s = resolve_subject(v, curriculum_codes | {"__"})
        if s:
            posts[j] = (v, s)
            if s not in curriculum_codes:
                warnings.append(("w_not_in_curr", {"p": v, "s": s}))

    # 2. teacher names row (اللقب و الاسم) – optional: first row below posts that has text in post columns
    names = {}
    for i in range(post_row + 1, min(post_row + 3, len(grid))):
        row = grid.iloc[i]
        if any(class_to_id(v) for v in row):
            break
        vals = {j: row.iloc[j] for j in posts if row.iloc[j] and not re.fullmatch(r"[\d.]+", row.iloc[j])}
        if vals:
            names = vals
            break

    # 3. class rows + remedial row
    records, class_hours, remedial = [], {}, {}
    for i in range(post_row + 1, len(grid)):
        row = grid.iloc[i]
        if any("استدراك" in v for v in row):
            for j, (p, _) in posts.items():
                if _num(row.iloc[j]) > 0:
                    remedial[names.get(j, p)] = _num(row.iloc[j])
            continue
        cid = None
        for j, v in enumerate(row):
            if j in posts:
                continue
            cid = class_to_id(v)
            if cid:
                cls_col = j
                break
        if not cid:
            continue
        # declared weekly hours = first number next to the class label that is outside post columns
        for j in (cls_col - 1, cls_col + 1):
            if 0 <= j < len(row) and j not in posts and _num(row.iloc[j]) > 0:
                class_hours[cid] = _num(row.iloc[j])
                break
        for j, (p, s) in posts.items():
            h = _num(row.iloc[j])
            if h > 0:
                records.append({"Teacher": names.get(j, p), "Post": p, "Class": cid, "Subject": s, "Hours": int(h)})

    if not records:
        raise ValueError(t("w_no_classes"))
    assign = pd.DataFrame(records)

    # 4. consistency checks
    dup = assign.groupby(["Class", "Subject"])["Teacher"].nunique()
    for (c, s), n in dup[dup > 1].items():
        warnings.append(("w_split", {"c": c, "s": s, "n": n,
                                     "t": assign[(assign.Class == c) & (assign.Subject == s)].iloc[0]['Teacher']}))
    assign = assign.drop_duplicates(["Class", "Subject"], keep="first").reset_index(drop=True)
    if known_classes is not None:
        unknown = sorted(set(assign["Class"]) - set(known_classes))
        if unknown:
            warnings.append(("w_unknown_cls", {"x": ", ".join(unknown)}))

    teachers = (assign.groupby("Teacher")
                .agg(Qualified_Subjects=("Subject", lambda x: ";".join(sorted(set(x)))), Hours=("Hours", "sum"))
                .reset_index())
    teachers["Max_Weekly_Hours"] = teachers["Hours"] + teachers["Teacher"].map(remedial).fillna(0).astype(int)
    teachers["Requires_Reception"] = False
    teachers = teachers.rename(columns={"Teacher": "Teacher_ID"})[
        ["Teacher_ID", "Qualified_Subjects", "Max_Weekly_Hours", "Requires_Reception"]]
    classes = pd.DataFrame({"Class_ID": sorted(assign["Class"].unique(), key=_class_key)})
    classes["Level"] = classes["Class_ID"].str.extract(r"^(\d+AM)")[0]
    return {"assignment": assign, "teachers": teachers, "classes": classes, "remedial": remedial,
            "class_hours": class_hours, "warnings": warnings}


def _class_key(c):
    m = re.match(r"(\d+)\D+(\d+)", c)
    return (-int(m.group(1)), int(m.group(2))) if m else (0, 0)


# ---------------------------------------------------------------- export
def post_names(assignment):
    """Map each (teacher, subject) pair to a post column label (one post = one subject, like the school grid).
    Keeps an existing 'Post' column, else builds عربية1, عربية2 …"""
    if "Post" in assignment.columns:
        return {(t, s): p for t, s, p in zip(assignment["Teacher"], assignment["Subject"], assignment["Post"])}
    out, counters = {}, {}
    for t, s in sorted(set(zip(assignment["Teacher"], assignment["Subject"]))):
        counters[s] = counters.get(s, 0) + 1
        out[(t, s)] = f"{CODE_TO_ARABIC.get(s, s)}{counters[s]}"
    return out


def build_matrix_xlsx(assignment, hours_fn, class_levels, remedial=None, arabic_class_names=True):
    """assignment: Teacher/Class/Subject[/Hours][/Post]; hours_fn(level, subject) -> weekly hours (if no Hours col)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    a = assignment.copy()
    if "Hours" not in a.columns:
        a["Hours"] = [hours_fn(class_levels[c], s) for c, s in zip(a["Class"], a["Subject"])]
    posts = post_names(a)
    subj_rank = {k: i for i, k in enumerate(CODE_TO_ARABIC)}
    order = sorted(posts, key=lambda ts: (subj_rank.get(ts[1], 99), _natural(posts[ts])))
    classes = sorted(a["Class"].unique(), key=_class_key)
    remedial = remedial or {}
    lab = (lambda c: re.sub(r"^(\d+)AM(\d+)$", r"\1م\2", c)) if arabic_class_names else (lambda c: c)
    # remedial hours belong to a teacher: show them on that teacher's first post only
    rem_col, seen = {}, set()
    for ts in order:
        t = ts[0]
        if t not in seen:
            seen.add(t); rem_col[ts] = remedial.get(t, remedial.get(posts[ts], 0))

    wb = Workbook(); ws = wb.active; ws.title = "توزيع الحصص"; ws.sheet_view.rightToLeft = True
    thin = Side(style="thin", color="999999"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFFF00"); bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rot = Alignment(text_rotation=90, horizontal="center", vertical="center")
    first = 4
    hours = {ts: int(a[(a.Teacher == ts[0]) & (a.Subject == ts[1])]["Hours"].sum()) for ts in order}
    totals = {ts: hours[ts] + int(rem_col.get(ts, 0)) for ts in order}

    ws.cell(1, 1, "الحجم الساعي"); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.cell(1, 3, sum(totals.values()))
    ws.cell(2, 1, "المناصب"); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=3)
    ws.cell(3, 1, "اللقب و الاسم"); ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=3)
    for k, h in enumerate(["رقم", "الافواج", "حجم ساعي"]):
        ws.cell(4, k + 1, h).font = bold
    for j, ts in enumerate(order):
        col = first + j
        ws.cell(1, col, totals[ts])
        c = ws.cell(2, col, posts[ts]); c.font = bold; c.alignment = rot
        if posts[ts] != ts[0]:
            ws.cell(3, col, ts[0]).alignment = rot
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 5
    r = 5
    for n, c in enumerate(classes, 1):
        row = a[a.Class == c]
        ws.cell(r, 1, n); ws.cell(r, 2, lab(c)).font = bold; ws.cell(r, 3, int(row["Hours"].sum()))
        for j, ts in enumerate(order):
            h = row[(row.Teacher == ts[0]) & (row.Subject == ts[1])]["Hours"].sum()
            if h:
                ws.cell(r, first + j, int(h))
        if n % 2:
            for col in range(1, first + len(order)):
                ws.cell(r, col).fill = yellow
        r += 1
    if any(rem_col.values()):
        ws.cell(r, 1, "الاستدراك").font = Font(bold=True, color="FF0000")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        for j, ts in enumerate(order):
            if rem_col.get(ts):
                ws.cell(r, first + j, int(rem_col[ts])).font = Font(color="FF0000", bold=True)
        r += 1
    ws.cell(r, 1, "المجموع").font = bold
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for j, ts in enumerate(order):
        ws.cell(r, first + j, totals[ts]).font = bold
    ws.row_dimensions[2].height = 80; ws.row_dimensions[3].height = 60
    ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 9
    for row in ws.iter_rows(min_row=1, max_row=r, max_col=first + len(order) - 1):
        for c in row:
            c.border = border
            if c.alignment.text_rotation == 0:
                c.alignment = center
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _natural(s):
    return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", str(s))]


def validate_against_curriculum(assign, class_levels, hours_fn, available_slots):
    """Compare matrix hours with curriculum-derived hours and class capacity."""
    rows = []
    for _, r in assign.iterrows():
        need = hours_fn(class_levels.get(r["Class"]), r["Subject"])
        if need != r["Hours"]:
            rows.append({"Class": r["Class"], "Subject": r["Subject"], "Teacher": r["Teacher"],
                         "Matrix hours": r["Hours"], "Curriculum hours": need})
    cap = assign.groupby("Class")["Hours"].sum()
    over = cap[cap > available_slots]
    return pd.DataFrame(rows), over
